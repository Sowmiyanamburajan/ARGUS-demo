from flask import Flask, request, render_template, jsonify, send_from_directory,send_file
from pathlib import Path
import uuid, os
import numpy as np
from werkzeug.utils import secure_filename
from src.fake_document_detection import detect_fake_document
from pdf2image import convert_from_path
from src.detectors.document_detector import convert_document_to_image
from docx import Document
from PIL import Image, ImageDraw
from docx import Document
from datetime import datetime, timedelta
# ---- Import project modules ----
from src.chatbot import chat_receive_dynamic, get_chat_history
from src.audio_emotion import analyze_audio_emotion, transcribe_audio
from src.detectors.image_detector import detect_image
from src.detectors.video_detector import detect_video
from src.detectors.audio_detector import detect_audio
from src.registry import append_registry_entry
from src.behavioral import save_sample, save_profile
from src.cross_modal import lip_sync_score, cross_modal_consensus
from src.provenance import check_sidecar_signature, check_blockchain_stub, comprehensive_provenance_check, generate_provenance_certificate
from src.voice_stress import voice_stress_score
from src.emotion_manipulation import flag_emotion_manipulation
from src.factcheck import check_claims
from src.voice_face_identity import check_voice_face_identity
from src.explainable_output import generate_explainable_output
from src.consistency_check import check_modality_consistency





app = Flask(__name__, static_folder="static", template_folder="templates")

# ---- Directories ----
UPLOAD_DIR = Path("uploads"); UPLOAD_DIR.mkdir(exist_ok=True)
OUT_DIR = Path("outputs"); OUT_DIR.mkdir(exist_ok=True)
MODEL_DIR = Path("models")
if not MODEL_DIR.exists():
    raise FileNotFoundError(f"DNN model files not found in {MODEL_DIR}")


# -------------------------
# BEHAVIORAL ENDPOINTS
# -------------------------
@app.route('/behavior/profile', methods=['POST'])
def behavior_profile():
    payload = request.get_json(force=True)
    pid = payload.get('profile_id', 'guest')
    save_profile(pid, payload.get('user', 'guest'),
                 payload.get('typing_stats', {}),
                 payload.get('mouse_stats', {}))
    return jsonify({'saved': True})

@app.route('/behavior/sample', methods=['POST'])
def behavior_sample():
    payload = request.get_json(force=True)
    pid = payload.get('profile_id', 'guest')
    save_sample(pid,
                payload.get('typing_stats', {}),
                payload.get('mouse_stats', {}))
    return jsonify({'saved': True})

# -------------------------
# CROSS-MODAL ENDPOINTS
# -------------------------
@app.route('/cross/lipsync', methods=['POST'])
def cross_lipsync():
    payload = request.get_json(force=True)
    frames_dir = payload.get('frames_dir')
    wav_path = payload.get('wav_path')
    res = lip_sync_score(frames_dir, wav_path)
    return jsonify(res)

# -------------------------
# PROVENANCE ENDPOINTS
# -------------------------
@app.route('/provenance/check', methods=['POST'])
def provenance_check():
    payload = request.get_json(force=True)
    fpath = payload.get('file_path')
    res = check_sidecar_signature(fpath)
    res2 = check_blockchain_stub(fpath)
    return jsonify({'sidecar': res, 'on_chain': res2})

# -------------------------
# CHAT ENDPOINTS
# -------------------------
@app.route("/")
def index():
    return render_template("index.html")

@app.route("/chat")
def chat_page():
    return render_template("chat.html")

@app.route("/chat/answer", methods=["POST"])
def chat_answer():
    data = request.json or {}
    session_id = data.get("session_id")
    text = data.get("text")
    result = chat_receive_dynamic(session_id=session_id, answer_text=text)
    return jsonify(result)

@app.route("/chat/answer_audio", methods=["POST"])
def chat_answer_audio():
    session_id = request.form.get("session_id")
    f = request.files.get("audio")
    if not f:
        return jsonify({"error": "No audio uploaded"}), 400

    filename = str(uuid.uuid4())[:8] + "_resp.wav"
    path = UPLOAD_DIR / filename
    f.save(path)

    transcript = None
    try:
        transcript = transcribe_audio(str(path))
    except Exception:
        transcript = None

    audio_em = analyze_audio_emotion(str(path))
    result = chat_receive_dynamic(
        session_id=session_id,
        answer_text=transcript or None,
        audio_path=str(path)
    )
    result["transcript"] = transcript
    result["audio_emotion"] = audio_em
    return jsonify(result)


# @app.route("/analyze", methods=["POST"])
# def analyze():
#     files = list(request.files.getlist("file"))
#     results = []

#     for f in files:
#         if not f:
#             continue

#         filename = str(uuid.uuid4())[:8] + "_" + secure_filename(f.filename)
#         dest = UPLOAD_DIR / filename
#         f.save(dest)

#         file_ext = f.filename.lower().split('.')[-1] if '.' in f.filename else ''
#         file_type = (
#             'image' if file_ext in ['jpg', 'jpeg', 'png', 'gif', 'bmp', 'tiff', 'webp'] else
#             'audio' if file_ext in ['wav', 'mp3', 'm4a', 'flac', 'aac', 'ogg', 'wma'] else
#             'video' if file_ext in ['mp4', 'avi', 'mov', 'mkv', 'wmv', 'flv', 'webm', 'm4v'] else
#             'document' if file_ext in ['pdf', 'doc', 'docx', 'txt', 'rtf'] else
#             'unknown'
#         )

#         basic_analysis = {'type': file_type, 'faces': []}

#         try:
#             if file_type == 'image':
#                 basic_analysis = detect_image(str(dest), str(OUT_DIR))
#                 basic_analysis['preview_image'] = f"/uploads/{filename}"
#                 basic_analysis['thumbnail'] = f"/uploads/{filename}"

#             elif file_type == 'video':
#                 basic_analysis = detect_video(str(dest), str(OUT_DIR))

#             elif file_type == 'audio':
#                 basic_analysis = detect_audio(str(dest), str(OUT_DIR))

#             elif file_type == 'document':
#                 img_path = None
#                 if file_ext == 'pdf':
#                     images = convert_from_path(dest, dpi=200, first_page=1, last_page=1)
#                     img_path = OUT_DIR / f"{filename}_page1.jpg"
#                     images[0].save(img_path, "JPEG")
#                 elif file_ext == 'docx':
#                     doc = Document(dest)
#                     text = "\n".join([p.text for p in doc.paragraphs])
#                     img = Image.new("RGB", (1000, 1200), color="white")
#                     from PIL import ImageDraw
#                     draw = ImageDraw.Draw(img)
#                     draw.text((20, 20), text[:2000], fill="black")
#                     img_path = OUT_DIR / f"{filename}_preview.jpg"
#                     img.save(img_path)

#                 if img_path and os.path.exists(img_path):
#                     doc_analysis = detect_fake_document(str(img_path))
#                     basic_analysis.update(doc_analysis)
#                     basic_analysis['preview_image'] = str(img_path)
#                     basic_analysis['thumbnail'] = str(img_path)
#                 else:
#                     basic_analysis['explanation'] = "Document could not be converted to image for analysis."
            
#         except Exception as e:
#             basic_analysis['error'] = f"Error analyzing {file_type}: {str(e)}"
#             print(f"Error in {file_type} detection: {e}")

#         enhanced_results = {}
#         try:
#             if file_type == 'image':
#                 enhanced_results['explainable_output'] = generate_explainable_output(
#                     str(dest), 'image', basic_analysis, str(OUT_DIR)
#                 )

#             elif file_type == 'audio':
#                 enhanced_results['audio_emotion'] = analyze_audio_emotion(str(dest))
#                 enhanced_results['voice_stress'] = voice_stress_score(str(dest))
#                 enhanced_results['explainable_output'] = generate_explainable_output(
#                     str(dest), 'audio', basic_analysis, str(OUT_DIR)
#                 )

#             elif file_type == 'video':
#                 enhanced_results['lip_sync'] = lip_sync_score(str(dest), str(dest))
#                 enhanced_results['explainable_output'] = generate_explainable_output(
#                     str(dest), 'video', basic_analysis, str(OUT_DIR)
#                 )

#             elif file_type == 'document' and basic_analysis.get('preview_image'):
#                 enhanced_results['explainable_output'] = generate_explainable_output(
#                     basic_analysis['preview_image'], 'document', basic_analysis, str(OUT_DIR)
#                 )

#         except Exception as e:
#             enhanced_results['error'] = str(e)
#             print(f"Error in enhanced results: {e}")

#         # Safe final result for template
#         final_result = {
#             'file_path': str(dest),
#             'file_type': file_type,
#             'filename': filename,
#             'preview_image': basic_analysis.get('preview_image'),
#             'thumbnail': basic_analysis.get('thumbnail'),
#             'prediction': basic_analysis.get('prediction', 'N/A'),
#             'confidence': basic_analysis.get('confidence', 0),
#             'metadata': basic_analysis.get('metadata'),
#             'score': basic_analysis.get('score', 0),
#             'ai_prob': basic_analysis.get('ai_prob', 0),
#             'text_tamper': basic_analysis.get('text_tamper'),
#             'metadata_tamper': basic_analysis.get('metadata_tamper'),
#             'heatmap_image': basic_analysis.get('heatmap_image'),
#             'heatmap_overlay': basic_analysis.get('heatmap_overlay'),
#             'heatmap_data': basic_analysis.get('heatmap_data'),
#             'ocr_text': basic_analysis.get('ocr_text'),
#             'sharpness': basic_analysis.get('sharpness'),
#             'high_freq_energy': basic_analysis.get('high_freq_energy'),
#             'faces': basic_analysis.get('faces', []),
#             'faces_boxed': basic_analysis.get('faces_boxed'),
#             'frame_scores': basic_analysis.get('frame_scores', []),
#             'frame_labels': basic_analysis.get('frame_labels', []),
#             'explanation': basic_analysis.get('explanation', 'No explanation available.'),
#             'summary': basic_analysis.get('summary', 'No summary available.'),
#             'explainable_output': enhanced_results.get('explainable_output'),
#             'error': basic_analysis.get('error') or enhanced_results.get('error')
#         }

#         results.append(final_result)

#     if results:
#         return render_template("result.html", result=results[0])
#     else:
#         return "No valid file uploaded", 400



POPPLER_PATH = r"C:\Users\jc user\Downloads\poppler-windows-25.07.0-0\Library\bin"

@app.route("/analyze", methods=["POST"])
def analyze():
    files = list(request.files.getlist("file"))
    results = []

    for f in files:
        if not f:
            continue

        filename = str(uuid.uuid4())[:8] + "_" + secure_filename(f.filename)
        dest = UPLOAD_DIR / filename
        UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
        f.save(dest)

        file_ext = f.filename.lower().split('.')[-1] if '.' in f.filename else ''
        file_type = (
            'image' if file_ext in ['jpg', 'jpeg', 'png', 'gif', 'bmp', 'tiff', 'webp'] else
            'audio' if file_ext in ['wav', 'mp3', 'm4a', 'flac', 'aac', 'ogg', 'wma'] else
            'video' if file_ext in ['mp4', 'avi', 'mov', 'mkv', 'wmv', 'flv', 'webm', 'm4v'] else
            'document' if file_ext in ['pdf', 'doc', 'docx', 'txt', 'rtf'] else
            'unknown'
        )

        basic_analysis = {'type': file_type, 'faces': []}

        try:
            if file_type == 'image':
                basic_analysis = detect_image(str(dest), str(OUT_DIR))
                basic_analysis['preview_image'] = f"/uploads/{filename}"
                basic_analysis['thumbnail'] = f"/uploads/{filename}"

            elif file_type == 'video':
                basic_analysis = detect_video(str(dest), str(OUT_DIR))

            elif file_type == 'audio':
                basic_analysis = detect_audio(str(dest), str(OUT_DIR))

            elif file_type == 'document':
                img_path = None
                try:
                    if file_ext == 'pdf':
                        images = convert_from_path(dest, dpi=200, first_page=1, last_page=1, poppler_path=POPPLER_PATH)
                        img_path = OUT_DIR / f"{filename}_page1.jpg"
                        OUT_DIR.mkdir(parents=True, exist_ok=True)
                        images[0].save(img_path, "JPEG")

                    elif file_ext == 'docx':
                        doc = Document(dest)
                        text = "\n".join([p.text for p in doc.paragraphs])
                        img = Image.new("RGB", (1000, 1200), color="white")
                        draw = ImageDraw.Draw(img)
                        draw.text((20, 20), text[:2000], fill="black")
                        img_path = OUT_DIR / f"{filename}_preview.jpg"
                        OUT_DIR.mkdir(parents=True, exist_ok=True)
                        img.save(img_path)

                    if img_path and os.path.exists(img_path):
                        doc_analysis = detect_fake_document(str(img_path))
                        basic_analysis.update(doc_analysis)
                        basic_analysis['preview_image'] = str(img_path)
                        basic_analysis['thumbnail'] = str(img_path)
                    else:
                        basic_analysis['explanation'] = "Document could not be converted to image for analysis."

                except Exception as e:
                    basic_analysis['explanation'] = f"PDF conversion failed: {e}"

        except Exception as e:
            basic_analysis['error'] = f"Error analyzing {file_type}: {str(e)}"
            print(f"Error in {file_type} detection: {e}")

        # Enhanced results (explainable outputs, emotions, lip-sync)
        enhanced_results = {}
        try:
            if file_type == 'image':
                enhanced_results['explainable_output'] = generate_explainable_output(
                    str(dest), 'image', basic_analysis, str(OUT_DIR)
                )

            elif file_type == 'audio':
                enhanced_results['audio_emotion'] = analyze_audio_emotion(str(dest))
                enhanced_results['voice_stress'] = voice_stress_score(str(dest))
                enhanced_results['explainable_output'] = generate_explainable_output(
                    str(dest), 'audio', basic_analysis, str(OUT_DIR)
                )

            elif file_type == 'video':
                enhanced_results['lip_sync'] = lip_sync_score(str(dest), str(dest))
                enhanced_results['explainable_output'] = generate_explainable_output(
                    str(dest), 'video', basic_analysis, str(OUT_DIR)
                )

            elif file_type == 'document' and basic_analysis.get('preview_image'):
                enhanced_results['explainable_output'] = generate_explainable_output(
                    basic_analysis['preview_image'], 'document', basic_analysis, str(OUT_DIR)
                )

        except Exception as e:
            enhanced_results['error'] = str(e)
            print(f"Error in enhanced results: {e}")

        # Safe final result for template
        final_result = {
            'file_path': str(dest),
            'file_type': file_type,
            'filename': filename,
            'preview_image': basic_analysis.get('preview_image'),
            'thumbnail': basic_analysis.get('thumbnail'),
            'prediction': basic_analysis.get('prediction', 'N/A'),
            'confidence': basic_analysis.get('confidence', 0),
            'metadata': basic_analysis.get('metadata'),
            'score': basic_analysis.get('score', 0),
            'ai_prob': basic_analysis.get('ai_prob', 0),
            'text_tamper': basic_analysis.get('text_tamper'),
            'metadata_tamper': basic_analysis.get('metadata_tamper'),
            'heatmap_image': basic_analysis.get('heatmap_image'),
            'heatmap_overlay': basic_analysis.get('heatmap_overlay'),
            'heatmap_data': basic_analysis.get('heatmap_data'),
            'ocr_text': basic_analysis.get('ocr_text'),
            'sharpness': basic_analysis.get('sharpness'),
            'high_freq_energy': basic_analysis.get('high_freq_energy'),
            'faces': basic_analysis.get('faces', []),
            'faces_boxed': basic_analysis.get('faces_boxed'),
            'frame_scores': basic_analysis.get('frame_scores', []),
            'frame_labels': basic_analysis.get('frame_labels', []),
            'explanation': basic_analysis.get('explanation', 'No explanation available.'),
            'summary': basic_analysis.get('summary', 'No summary available.'),
            'explainable_output': enhanced_results.get('explainable_output'),
            'error': basic_analysis.get('error') or enhanced_results.get('error')
        }

        results.append(final_result)

    if results:
        return render_template("result.html", result=results[0])
    else:
        return "No valid file uploaded", 400




@app.route("/download_ocr/<filename>")
def download_ocr(filename):
    ocr_file_path = OCR_DIR / f"{filename}_ocr.txt"
    if os.path.exists(ocr_file_path):
        return send_file(ocr_file_path, as_attachment=True)
    return "OCR file not available", 404

# -------------------------
# HISTORY ENDPOINT
# -------------------------
@app.route("/chat/history/<session_id>", methods=["GET"])
def chat_history(session_id):
    hist = get_chat_history(session_id)
    return jsonify(hist)

# -------------------------
# STATIC FILES (OUTPUTS)
# -------------------------
@app.route('/outputs/<path:filename>')
def outputs(filename):
    return send_from_directory('outputs', filename)

@app.route('/uploads/<path:filename>')
def uploaded_file(filename):
    return send_from_directory('uploads', filename)

@app.route("/image-detection")
def image_detection():
    return render_template("image.html")

@app.route("/video-detection")
def video_detection():
    return render_template("video.html")

@app.route("/audio-detection")
def audio_detection():
    return render_template("audio.html")

# -------------------------
# ADVANCED FEATURES ENDPOINTS
# -------------------------

@app.route('/advanced/voice-face-identity', methods=['POST'])
def voice_face_identity_check():
    """Cross-check voice and face identity verification"""
    try:
        data = request.get_json()
        audio_path = data.get('audio_path')
        image_path = data.get('image_path')
        known_identities = data.get('known_identities', {})
        
        result = check_voice_face_identity(audio_path, image_path, known_identities)
        return jsonify(result)
    except Exception as e:
        return jsonify({'error': str(e)}), 400

@app.route('/advanced/emotion-manipulation', methods=['POST'])
def emotion_manipulation_check():
    """Enhanced emotion manipulation detection"""
    try:
        data = request.get_json()
        text_emotion = data.get('text_emotion')
        audio_emotion = data.get('audio_emotion')
        video_emotion = data.get('video_emotion')
        audio_path = data.get('audio_path')
        image_path = data.get('image_path')
        
        result = flag_emotion_manipulation(
            text_emotion, audio_emotion, video_emotion,
            audio_path, image_path
        )
        return jsonify(result)
    except Exception as e:
        return jsonify({'error': str(e)}), 400

@app.route('/advanced/provenance-check', methods=['POST'])
def comprehensive_provenance_check_endpoint():
    """Comprehensive provenance verification"""
    try:
        data = request.get_json()
        file_path = data.get('file_path')
        
        result = comprehensive_provenance_check(file_path)
        return jsonify(result)
    except Exception as e:
        return jsonify({'error': str(e)}), 400

@app.route('/advanced/generate-certificate', methods=['POST'])
def generate_certificate():
    """Generate provenance certificate for a file"""
    try:
        data = request.get_json()
        file_path = data.get('file_path')
        issuer = data.get('issuer', 'ARGUS_System')
        
        certificate = generate_provenance_certificate(file_path, issuer)
        return jsonify(certificate)
    except Exception as e:
        return jsonify({'error': str(e)}), 400

@app.route('/advanced/explainable-output', methods=['POST'])
def explainable_output_generation():
    """Generate explainable output with heatmaps and analysis"""
    try:
        data = request.get_json()
        file_path = data.get('file_path')
        file_type = data.get('file_type')  # 'image', 'audio', 'video'
        detection_results = data.get('detection_results', {})
        
        result = generate_explainable_output(file_path, file_type, detection_results, str(OUT_DIR))
        return jsonify(result)
    except Exception as e:
        return jsonify({'error': str(e)}), 400

@app.route('/advanced/consistency-check', methods=['POST'])
def consistency_check_endpoint():
    """Check consistency across modalities"""
    try:
        data = request.get_json()
        text = data.get('text')
        media_path = data.get('media_path')
        media_type = data.get('media_type')  # 'image', 'audio', 'video'
        
        result = check_modality_consistency(text, media_path, media_type)
        return jsonify(result)
    except Exception as e:
        return jsonify({'error': str(e)}), 400

@app.route('/advanced/fake-document-detection', methods=['POST'])
def fake_document_detection():
    """Detect fake documents"""
    try:
        data = request.get_json()
        image_path = data.get('image_path')
        
        result = detect_fake_document(image_path)
        return jsonify(result)
    except Exception as e:
        return jsonify({'error': str(e)}), 400




if __name__ == "__main__":
    app.run(debug=True)

